"""Excel 描述清单：一次校验、撤销开放字段、只重算变化文本。"""

import pytest
from openpyxl import Workbook

from dm_core.catalog import (
    build_catalog,
    dataset_id,
    description_path,
    ontology_path,
    source_material,
)
from dm_core.config import AppError, read_json
from dm_core.index import sync
from dm_core.spreadsheet import (
    import_description_workbook,
    persist_description_workbook,
    read_workbook,
)


def _write_workbook(path, *, duplicate_field=False, name_description="事件名称"):
    workbook = Workbook()
    tables = workbook.active
    tables.title = "Tables"
    tables.append(
        [
            "schema",
            "table",
            "name",
            "description",
            "aliases",
            "grain",
            "default_time",
            "entity_id",
        ]
    )
    tables.append(
        ["S", "EVENT", "事件", "业务事件记录", "记录,活动", "一行一个事件", "AT", "event"]
    )
    tables.append(["S", "DEVICE", "设备", "设备主档", "", "", "", "device"])

    entities = workbook.create_sheet("Entities")
    entities.append(["entity_id", "name", "description", "aliases", "grain", "status"])
    entities.append(["event", "事件", "发生过的业务事件", "记录,活动", "一行一个事件", "verified"])
    entities.append(["device", "设备", "设备业务实体", "机器", "一台设备", "verified"])

    fields = workbook.create_sheet("Fields")
    fields.append(
        [
            "schema",
            "table",
            "field",
            "name",
            "description",
            "unit",
            "enum",
            "null_meaning",
            "entity_id",
            "property_id",
            "semantic_type",
            "semantic_role",
            "synonyms",
        ]
    )
    fields.append(
        [
            "S",
            "EVENT",
            "NAME",
            "名称",
            name_description,
            "",
            "",
            "",
            "event",
            "name",
            "name",
            "dimension",
            "标题,事件名",
        ]
    )
    fields.append(
        [
            "S",
            "EVENT",
            "AT",
            "",
            "事件发生时间",
            "",
            "",
            "",
            "event",
            "occurred_at",
            "datetime",
            "actual_time",
            "发生时刻",
        ]
    )
    fields.append(
        [
            "S",
            "EVENT",
            "DEVICE_ID",
            "设备编号",
            "关联设备的编号",
            "",
            "",
            "",
            "event",
            "device_reference",
            "identifier",
            "dimension",
            "",
        ]
    )
    if duplicate_field:
        fields.append(
            ["S", "EVENT", "NAME", "重复名称", "错误重复字段", "", "", "", "", "", "", "", ""]
        )
    fields.append(
        [
            "S",
            "DEVICE",
            "ID",
            "编号",
            "设备编号",
            "",
            "",
            "",
            "device",
            "record_id",
            "identifier",
            "identifier",
            "",
        ]
    )
    fields.append(
        [
            "S",
            "DEVICE",
            "NAME",
            "名称",
            "设备名称",
            "",
            "",
            "",
            "device",
            "name",
            "name",
            "dimension",
            "",
        ]
    )

    relations = workbook.create_sheet("Relations")
    relations.append(
        [
            "relation_id",
            "source_entity",
            "target_entity",
            "predicate",
            "cardinality",
            "description",
            "status",
        ]
    )
    relations.append(
        [
            "event_owned_by_device",
            "event",
            "device",
            "owned_by",
            "many_to_one",
            "事件所属设备",
            "verified",
        ]
    )

    paths = workbook.create_sheet("JoinPaths")
    paths.append(
        [
            "relation_id",
            "path_id",
            "schema",
            "table",
            "local_field",
            "target",
            "target_field",
            "description",
        ]
    )
    paths.append(
        [
            "event_owned_by_device",
            "device_pk",
            "S",
            "EVENT",
            "DEVICE_ID",
            "S.DEVICE",
            "ID",
            "事件外键到设备主键",
        ]
    )

    metrics = workbook.create_sheet("Metrics")
    metrics.append(["schema", "table", "id", "name", "description", "aggregation", "field", "unit"])
    metrics.append(["S", "EVENT", "record_count", "记录数", "本表记录数量", "count", "NAME", "条"])
    workbook.save(path)


def test_workbook_import_builds_tables_fields_metrics_and_relations(settings, metadata, tmp_path):
    path = tmp_path / "catalog.xlsx"
    _write_workbook(path)
    summary = import_description_workbook(settings, path)
    assert summary["tables"] == 2 and summary["fields"] == 5

    raw = read_json(description_path(settings, "S.EVENT"))
    assert raw["aliases"] == ["记录", "活动"]
    assert raw["default_time"] == "AT"
    assert raw["fields"]["NAME"]["name"] == "名称"
    assert raw["fields"]["AT"]["description"] == "事件发生时间"
    assert raw["fields"]["AT"]["semantic_role"] == "actual_time"
    assert raw["metrics"][0]["id"] == "record_count"
    ontology = read_json(ontology_path(settings))
    assert ontology["entities"][0]["id"] == "event"
    assert ontology["relations"][0]["path"]["on"] == {"DEVICE_ID": "ID"}

    catalog, _ = build_catalog(settings, metadata)
    event = next(iter(catalog.datasets.values()))
    assert event.name == "事件"
    assert len(event.fields) == 3
    assert event.relations[0].description == "事件所属设备"
    assert event.entity_id == "event"
    assert catalog.entities["event"].name == "事件"
    assert catalog.context([event.id])[0]["ontology"]["entities"]


def test_invalid_workbook_does_not_overwrite_existing_config(settings, tmp_path):
    good = tmp_path / "good.xlsx"
    bad = tmp_path / "bad.xlsx"
    _write_workbook(good)
    import_description_workbook(settings, good)
    before_tables = (settings.root / "configs/tables.json").read_bytes()
    before_name = description_path(settings, "S.EVENT").read_bytes()

    _write_workbook(bad, duplicate_field=True)
    with pytest.raises(AppError, match="重复声明"):
        import_description_workbook(settings, bad)
    assert (settings.root / "configs/tables.json").read_bytes() == before_tables
    assert description_path(settings, "S.EVENT").read_bytes() == before_name


def test_workbook_can_be_built_in_memory_before_publish(settings, metadata, tmp_path):
    path = tmp_path / "catalog.xlsx"
    _write_workbook(path)
    before_tables = (settings.root / "configs/tables.json").read_bytes()
    before_event = description_path(settings, "S.EVENT").read_bytes()

    spec = read_workbook(path)
    catalog, in_memory_hash = build_catalog(
        settings, metadata, descriptions=spec.descriptions, ontology=spec.ontology
    )
    assert len(catalog.datasets) == 2
    assert (settings.root / "configs/tables.json").read_bytes() == before_tables
    assert description_path(settings, "S.EVENT").read_bytes() == before_event

    persist_description_workbook(settings, spec.descriptions, spec.ontology)
    _, _, published_hash = source_material(settings)
    assert published_hash == in_memory_hash


def test_pending_semantic_relation_is_context_only(settings, metadata, tmp_path):
    path = tmp_path / "catalog.xlsx"
    _write_workbook(path)
    spec = read_workbook(path)
    spec.ontology["relations"][0]["status"] = "pending"
    spec.ontology["relations"][0]["path"] = None

    catalog, _ = build_catalog(
        settings, metadata, descriptions=spec.descriptions, ontology=spec.ontology
    )
    event = catalog.get(dataset_id("S.EVENT"))
    assert event.relations == []
    assert next(iter(catalog.semantic_relations.values())).status == "pending"
    assert catalog.context([event.id])[0]["ontology"]["semantic_relations"]


def test_verified_semantic_relation_requires_target_primary_key(settings, metadata, tmp_path):
    path = tmp_path / "catalog.xlsx"
    _write_workbook(path)
    spec = read_workbook(path)
    spec.ontology["relations"][0]["path"]["on"] = {"DEVICE_ID": "NAME"}
    with pytest.raises(AppError, match="主键"):
        build_catalog(settings, metadata, descriptions=spec.descriptions, ontology=spec.ontology)


def test_workbook_sync_embeds_only_changed_description(settings, metadata, embeddings, tmp_path):
    path = tmp_path / "catalog.xlsx"
    _write_workbook(path)
    import_description_workbook(settings, path)
    catalog, source = build_catalog(settings, metadata)
    sync(settings, catalog, source, embeddings)

    embeddings.calls.clear()
    import_description_workbook(settings, path)
    unchanged, source = build_catalog(settings, metadata)
    assert sync(settings, unchanged, source, embeddings)["action"] == "reused"
    assert embeddings.calls == []

    _write_workbook(path, name_description="事件名称，也称活动标题")
    import_description_workbook(settings, path)
    changed, source = build_catalog(settings, metadata)
    result = sync(settings, changed, source, embeddings)
    assert result["embedded"] == 1
    assert len(embeddings.calls) == 1
